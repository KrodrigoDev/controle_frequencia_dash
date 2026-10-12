import io

from openpyxl import load_workbook

from src.model import crud
from src.services import exportar_excel as xls
from src.services import exportar_pdf as pdf
from src.services import indicadores as ind
from src.services import relatorios as rel


def test_excel_fechamento(meses):
    pac = rel.fechamento(ind.periodo_mes(*meses[-1]))
    wb = load_workbook(io.BytesIO(xls.excel_fechamento(pac)))
    assert wb.sheetnames[:3] == ["Resumo", "Por servidor", "Quadro do mês"]
    assert {"Ocorrências", "Pendências", "Lançamentos"} <= set(wb.sheetnames)
    assert wb["Lançamentos"].max_row == pac.resumo["registros"] + 1


def test_pdfs(meses):
    periodo = ind.periodo_mes(*meses[-1])
    assert pdf.pdf_fechamento(rel.fechamento(periodo)).startswith(b"%PDF")
    s = crud.list_servidores().iloc[0].to_dict()
    assert pdf.pdf_individual([(rel.individual(periodo, int(s["id"])), s)]).startswith(b"%PDF")
    assert pdf.pdf_individual([(rel.individual(ind.periodo_ano(periodo.ano), int(s["id"])), s)]).startswith(b"%PDF")


def test_periodo_vazio_nao_quebra():
    pac = rel.fechamento(ind.periodo_mes(2001, 1))
    assert pac.resumo["registros"] == 0
    assert xls.excel_fechamento(pac)[:2] == b"PK"
    assert pdf.pdf_fechamento(pac).startswith(b"%PDF")
